"""Auto-extracted from server.py."""
import uuid
import logging
from datetime import datetime, timezone, timedelta, date
from typing import Optional, List
from fastapi import APIRouter, HTTPException, Depends, UploadFile, File, Body, Request
from pydantic import BaseModel

router = APIRouter()

from core import (
    sdb, supabase_admin, supabase_user, logger, now_iso,
    generate_password, generate_slug, enrich_member, get_auth_users_map,
    get_plan_limits, parse_presentations,
    require_clinic_member, require_clinic_admin, require_super_admin, get_current_user,
    LoginRequest, LoginResponse, ClinicCreate, ClinicUpdate, ClinicMemberCreate, UserUpdate,
    MedicationCreate, MedicationBulkImport, LabStudyCreate, LabStudyBulkImport,
    ICD10CodeCreate, ICD10BulkImport,
    AppointmentCreate, AppointmentUpdate, AppointmentStatusUpdate,
    PatientQuickCreate, PatientFullCreate,
    validate_uuid, require_module,
)

# ============== PATIENT ROUTES (CLINIC) ==============

@router.get("/clinic/patients/search")
async def search_patients(q: str = "", ctx=Depends(require_clinic_member)):
    clinic_id = ctx["member"]["clinic_id"]
    try:
        if not q or len(q) < 2:
            result = sdb.table('patients').select('id,first_name,last_name,phone,national_id').eq('clinic_id', clinic_id).eq('is_active', True).order('first_name').limit(20).execute()
        else:
            from services.input_sanitizer import sanitize_postgrest_search
            qs = sanitize_postgrest_search(q)
            result = sdb.table('patients').select('id,first_name,last_name,phone,national_id').eq('clinic_id', clinic_id).eq('is_active', True).or_(f'first_name.ilike.%{qs}%,last_name.ilike.%{qs}%,national_id.ilike.%{qs}%,phone.ilike.%{qs}%').limit(20).execute()
        return result.data or []
    except Exception as e:
        logger.error(f"Search patients error: {e}")
        raise HTTPException(status_code=500, detail="Error al buscar pacientes")

@router.get("/clinic/patients")
async def list_patients(
    q: str = "",
    status: Optional[str] = None,
    page: int = 1,
    limit: int = 25,
    ctx=Depends(require_module('patients'))
):
    clinic_id = ctx["member"]["clinic_id"]
    try:
        query = sdb.table('patients').select('id,first_name,last_name,phone,email,national_id,date_of_birth,gender,is_active,created_at', count='exact').eq('clinic_id', clinic_id)

        if q:
            from services.input_sanitizer import sanitize_postgrest_search
            qs = sanitize_postgrest_search(q)
            query = query.or_(f'first_name.ilike.%{qs}%,last_name.ilike.%{qs}%,national_id.ilike.%{qs}%,phone.ilike.%{qs}%')
        if status == 'active':
            query = query.eq('is_active', True)
        elif status == 'inactive':
            query = query.eq('is_active', False)

        offset = (page - 1) * limit
        result = query.order('first_name').range(offset, offset + limit - 1).execute()
        patients = result.data or []

        # Enrich with visit counts and last visit
        for p in patients:
            apts = sdb.table('appointments').select('starts_at', count='exact').eq('patient_id', p['id']).eq('status', 'completed').execute()
            p['visit_count'] = apts.count or 0
            last = sdb.table('appointments').select('starts_at').eq('patient_id', p['id']).neq('status', 'cancelled').order('starts_at', desc=True).limit(1).execute()
            p['last_visit'] = last.data[0]['starts_at'] if last.data else None

        return {
            "patients": patients,
            "total": result.count or 0,
            "page": page,
            "limit": limit,
            "pages": ((result.count or 0) + limit - 1) // limit,
        }
    except Exception as e:
        logger.error(f"List patients error: {e}")
        raise HTTPException(status_code=500, detail="Error al listar pacientes")

@router.post("/clinic/patients")
async def create_patient(data: PatientFullCreate, ctx=Depends(require_clinic_member)):
    clinic_id = ctx["member"]["clinic_id"]
    # Enforce plan patient limit (defense against plan-limit bypass)
    _clinic = sdb.table('clinics').select('plan,max_patients').eq('id', clinic_id).maybe_single().execute()
    _cdata = getattr(_clinic, 'data', None) or {}
    _limit = _cdata.get('max_patients') or get_plan_limits(_cdata.get('plan', 'free')).get('max_patients')
    if _limit:
        _count = sdb.table('patients').select('id', count='exact').eq('clinic_id', clinic_id).eq('is_active', True).execute()
        if (_count.count or 0) >= _limit:
            raise HTTPException(status_code=400, detail="Límite de pacientes alcanzado para su plan. Contacte al administrador.")
    try:
        now = now_iso()
        patient_id = str(uuid.uuid4())
        doc = {
            "id": patient_id,
            "clinic_id": clinic_id,
            "first_name": data.first_name,
            "last_name": data.last_name,
            "is_active": True,
            "created_at": now,
            "updated_at": now,
        }
        # Optional fields
        optional = ['date_of_birth','gender','national_id','nationality','phone','phone_secondary',
                     'email','address','city','state','country',
                     'emergency_contact_name','emergency_contact_relation','emergency_contact_phone',
                     'blood_type','allergies','chronic_conditions','current_medications',
                     'insurance_provider','insurance_policy_number','insurance_expiry','notes']
        for field in optional:
            val = getattr(data, field, None)
            if val is not None:
                doc[field] = val

        sdb.table('patients').insert(doc).execute()
        return {"id": patient_id, "first_name": data.first_name, "last_name": data.last_name, "message": "Paciente creado exitosamente"}
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Create patient error: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail="Error al crear paciente")

@router.get("/clinic/patients/{patient_id}")
async def get_patient(patient_id: str, ctx=Depends(require_module('patients'))):
    validate_uuid(patient_id, "patient_id")
    clinic_id = ctx["member"]["clinic_id"]
    try:
        result = sdb.table('patients').select('*').eq('id', patient_id).eq('clinic_id', clinic_id).maybe_single().execute()
        data = getattr(result, 'data', None) if result else None
        if not data:
            raise HTTPException(status_code=404, detail="Paciente no encontrado")

        patient = data
        # Remove search_vector from response
        patient.pop('search_vector', None)

        # Get visit stats
        completed = sdb.table('appointments').select('id', count='exact').eq('patient_id', patient_id).eq('status', 'completed').execute()
        total_apts = sdb.table('appointments').select('id', count='exact').eq('patient_id', patient_id).neq('status', 'cancelled').execute()
        last = sdb.table('appointments').select('starts_at,reason,doctor_id').eq('patient_id', patient_id).neq('status', 'cancelled').order('starts_at', desc=True).limit(1).execute()
        next_apt = sdb.table('appointments').select('starts_at,reason,doctor_id').eq('patient_id', patient_id).eq('status', 'scheduled').gte('starts_at', now_iso()).order('starts_at').limit(1).execute()

        # Get recent appointments
        recent = sdb.table('appointments').select('*').eq('patient_id', patient_id).order('starts_at', desc=True).limit(20).execute()
        for apt in (recent.data or []):
            d = sdb.table('clinic_members').select('first_name,last_name').eq('id', apt.get('doctor_id','')).maybe_single().execute()
            apt['doctor_name'] = f"{d.data['first_name']} {d.data['last_name']}" if d.data else ""

        # Get files
        files = []
        try:
            file_list = supabase_admin.storage.from_('patient-files').list(f"{clinic_id}/{patient_id}")
            for f in (file_list or []):
                if f.get('name'):
                    files.append({
                        "name": f['name'],
                        "size": f.get('metadata', {}).get('size', 0) if f.get('metadata') else 0,
                        "created_at": f.get('created_at', ''),
                        "content_type": f.get('metadata', {}).get('mimetype', '') if f.get('metadata') else '',
                    })
        except Exception:
            pass

        return {
            "patient": patient,
            "stats": {
                "completed_visits": completed.count or 0,
                "total_appointments": total_apts.count or 0,
                "last_visit": last.data[0] if last.data else None,
                "next_appointment": next_apt.data[0] if next_apt.data else None,
            },
            "appointments": recent.data or [],
            "files": files,
        }
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Get patient error: {e}")
        raise HTTPException(status_code=500, detail="Error al obtener paciente")

@router.put("/clinic/patients/{patient_id}")
async def update_patient(patient_id: str, data: PatientFullCreate, ctx=Depends(require_module('patients'))):
    clinic_id = ctx["member"]["clinic_id"]
    try:
        existing = sdb.table('patients').select('id').eq('id', patient_id).eq('clinic_id', clinic_id).maybe_single().execute()
        if not existing.data:
            raise HTTPException(status_code=404, detail="Paciente no encontrado")

        update_data = {"updated_at": now_iso()}
        for field, val in data.model_dump().items():
            if val is not None:
                update_data[field] = val

        sdb.table('patients').update(update_data).eq('id', patient_id).execute()
        updated = sdb.table('patients').select('*').eq('id', patient_id).single().execute()
        updated.data.pop('search_vector', None)
        return updated.data
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Update patient error: {e}")
        raise HTTPException(status_code=500, detail="Error al actualizar paciente")

@router.put("/clinic/patients/{patient_id}/toggle-active")
async def toggle_patient_active(patient_id: str, ctx=Depends(require_module('patients'))):
    clinic_id = ctx["member"]["clinic_id"]
    try:
        patient = sdb.table('patients').select('is_active,first_name,last_name').eq('id', patient_id).eq('clinic_id', clinic_id).maybe_single().execute()
        if not patient.data:
            raise HTTPException(status_code=404, detail="Paciente no encontrado")
        new_val = not patient.data['is_active']
        sdb.table('patients').update({"is_active": new_val, "updated_at": now_iso()}).eq('id', patient_id).execute()
        try:
            from services.audit import log_audit, actor_from_ctx
            await log_audit(
                action="patient_activated" if new_val else "patient_archived",
                entity="patient",
                entity_id=patient_id,
                **actor_from_ctx(ctx),
                meta={"name": f"{patient.data.get('first_name','')} {patient.data.get('last_name','')}".strip()},
            )
        except Exception:
            pass
        return {"is_active": new_val}
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Toggle patient error: {e}")
        raise HTTPException(status_code=500, detail="Error")



# ============== PATIENT FILE ROUTES ==============

@router.post("/clinic/patients/{patient_id}/files")
async def upload_patient_file(patient_id: str, file: UploadFile = File(...), ctx=Depends(require_module('patients'))):
    clinic_id = ctx["member"]["clinic_id"]
    try:
        # Validate patient
        patient = sdb.table('patients').select('id').eq('id', patient_id).eq('clinic_id', clinic_id).maybe_single().execute()
        if not patient.data:
            raise HTTPException(status_code=404, detail="Paciente no encontrado")

        content = await file.read()
        # Magic-bytes validation (defeats Content-Type spoofing)
        try:
            from services.input_sanitizer import validate_document_upload
            detected_mime = validate_document_upload(
                content,
                max_bytes=10 * 1024 * 1024,
                extra_allowed={'application/dicom'},
            )
        except ValueError as ve:
            raise HTTPException(status_code=400, detail=str(ve))

        # Sanitize filename to avoid path traversal in storage paths
        import os as _os
        safe_name = _os.path.basename(file.filename or 'upload').replace('..', '').replace('/', '_').replace('\\', '_')[:120]
        # Upload to Supabase Storage
        path = f"{clinic_id}/{patient_id}/{safe_name}"
        supabase_admin.storage.from_('patient-files').upload(path, content, {"content-type": detected_mime, "upsert": "true"})

        return {"message": "Archivo subido", "name": safe_name, "size": len(content), "content_type": detected_mime}
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Upload file error: {e}")
        raise HTTPException(status_code=500, detail=f"Error al subir archivo: {str(e)}")

@router.get("/clinic/patients/{patient_id}/files/{filename}/url")
async def get_file_url(patient_id: str, filename: str, ctx=Depends(require_module('patients'))):
    clinic_id = ctx["member"]["clinic_id"]
    try:
        path = f"{clinic_id}/{patient_id}/{filename}"
        signed = supabase_admin.storage.from_('patient-files').create_signed_url(path, 3600)
        return {"url": signed.get('signedURL') or signed.get('signedUrl', '')}
    except Exception as e:
        logger.error(f"Get file URL error: {e}")
        raise HTTPException(status_code=500, detail="Error al obtener URL")

@router.delete("/clinic/patients/{patient_id}/files/{filename}")
async def delete_patient_file(patient_id: str, filename: str, ctx=Depends(require_module('patients'))):
    clinic_id = ctx["member"]["clinic_id"]
    try:
        path = f"{clinic_id}/{patient_id}/{filename}"
        supabase_admin.storage.from_('patient-files').remove([path])
        return {"message": "Archivo eliminado"}
    except Exception as e:
        logger.error(f"Delete file error: {e}")
        raise HTTPException(status_code=500, detail="Error al eliminar archivo")



# ============== BULK IMPORT (CSV / XLSX) ==============

PATIENT_REQUIRED = ["first_name", "last_name"]
PATIENT_OPTIONAL = [
    "date_of_birth", "gender", "national_id", "nationality",
    "phone", "phone_secondary", "email",
    "address", "city", "state", "country",
    "emergency_contact_name", "emergency_contact_relation", "emergency_contact_phone",
    "blood_type", "insurance_provider", "insurance_policy_number", "insurance_expiry",
    "notes",
]

def _parse_patient_row(row: dict) -> dict:
    """Normalize a row dict into a Supabase-ready patient document."""
    from datetime import datetime as _dt, date as _date
    def s(field):
        v = row.get(field)
        if v is None:
            return ""
        if isinstance(v, (_dt, _date)):
            return v.strftime("%Y-%m-%d")
        if isinstance(v, (int, float)):
            return str(v).strip()
        return str(v).strip()

    # Gender normalization
    g = s("gender").lower()
    gender = None
    if g in ("m", "male", "masculino", "hombre"):
        gender = "male"
    elif g in ("f", "female", "femenino", "mujer"):
        gender = "female"
    elif g in ("o", "other", "otro", "no_binary", "non_binary"):
        gender = "other"

    # Date of birth: accept ISO YYYY-MM-DD or DD/MM/YYYY
    dob = s("date_of_birth")
    if dob:
        try:
            from datetime import datetime as _dt
            if "/" in dob:
                dob = _dt.strptime(dob, "%d/%m/%Y").date().isoformat()
            else:
                dob = _dt.fromisoformat(dob[:10]).date().isoformat()
        except Exception:
            dob = ""

    return {
        "first_name": s("first_name"),
        "last_name": s("last_name"),
        "date_of_birth": dob or None,
        "gender": gender,
        "national_id": s("national_id") or None,
        "nationality": s("nationality") or None,
        "phone": s("phone") or None,
        "phone_secondary": s("phone_secondary") or None,
        "email": s("email") or None,
        "address": s("address") or None,
        "city": s("city") or None,
        "state": s("state") or None,
        "country": s("country") or None,
        "emergency_contact_name": s("emergency_contact_name") or None,
        "emergency_contact_relation": s("emergency_contact_relation") or None,
        "emergency_contact_phone": s("emergency_contact_phone") or None,
        "blood_type": s("blood_type") or None,
        "insurance_provider": s("insurance_provider") or None,
        "insurance_policy_number": s("insurance_policy_number") or None,
        "insurance_expiry": s("insurance_expiry") or None,
        "notes": s("notes") or None,
    }

def _parse_xlsx(raw: bytes) -> list:
    """Parse an .xlsx file into a list of row dicts (first sheet, first row = headers)."""
    import io as _io
    from openpyxl import load_workbook
    wb = load_workbook(filename=_io.BytesIO(raw), read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    headers = next(rows, None)
    if not headers:
        raise HTTPException(status_code=400, detail="XLSX sin encabezados")
    norm_headers = [(str(h or "").strip().lower().replace(" ", "_")) for h in headers]
    out = []
    for row in rows:
        if all(c is None or str(c).strip() == "" for c in row):
            continue
        out.append({norm_headers[i]: c for i, c in enumerate(row) if i < len(norm_headers)})
    return out

def _parse_csv(raw: bytes) -> list:
    """Parse CSV bytes into a list of row dicts with normalized headers."""
    import csv as _csv
    import io as _io
    text = None
    for enc in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise HTTPException(status_code=400, detail="No se pudo decodificar el CSV (use UTF-8)")
    sample = text[:2048]
    try:
        dialect = _csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except _csv.Error:
        dialect = _csv.excel
    reader = _csv.DictReader(_io.StringIO(text), dialect=dialect)
    if not reader.fieldnames:
        raise HTTPException(status_code=400, detail="CSV sin encabezados")
    norm_headers = {h: (h or "").strip().lower().replace(" ", "_") for h in reader.fieldnames}
    out = []
    for r in reader:
        d = {norm_headers[k]: v for k, v in r.items() if k in norm_headers}
        # Skip phantom rows from trailing newlines / fully blank lines
        if not any((str(v).strip() if v is not None else "") for v in d.values()):
            continue
        out.append(d)
    return out

@router.post("/clinic/patients-bulk/import")
async def import_patients(
    file: UploadFile = File(...),
    commit: bool = False,
    ctx=Depends(require_clinic_admin),
):
    """Bulk-import patients from a CSV or XLSX file.

    Required columns: first_name, last_name
    Dedup: by national_id (if provided) → fallback to first_name+last_name+phone
    """
    clinic_id = ctx["member"]["clinic_id"]
    raw = await file.read()
    if not raw:
        raise HTTPException(status_code=400, detail="Archivo vacío")
    if len(raw) > 5 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="Archivo supera el límite de 5 MB")
    fname = (file.filename or "").lower()
    if fname.endswith(".xlsx"):
        rows = _parse_xlsx(raw)
    elif fname.endswith(".csv"):
        rows = _parse_csv(raw)
    else:
        raise HTTPException(status_code=400, detail="Formato no soportado (use .csv o .xlsx)")

    if not rows:
        raise HTTPException(status_code=400, detail="Archivo sin filas de datos")

    # Verify required headers exist
    sample_keys = set(rows[0].keys())
    missing = [r for r in PATIENT_REQUIRED if r not in sample_keys]
    if missing:
        raise HTTPException(status_code=400, detail=f"Faltan columnas requeridas: {', '.join(missing)}")

    # Pre-fetch existing patient identifiers (national_id + name+phone composite) for dedup
    existing = sdb.table('patients').select('national_id,first_name,last_name,phone').eq('clinic_id', clinic_id).execute()
    existing_nat = set()
    existing_composite = set()
    for p in (existing.data or []):
        if p.get('national_id'):
            existing_nat.add(p['national_id'].strip().lower())
        comp = f"{(p.get('first_name') or '').strip().lower()}|{(p.get('last_name') or '').strip().lower()}|{(p.get('phone') or '').strip()}"
        existing_composite.add(comp)

    valid_rows = []
    errors = []
    skipped_dup = 0
    preview = []

    for idx, raw_row in enumerate(rows, start=2):  # row 1 = header
        try:
            doc = _parse_patient_row(raw_row)
            # Required validation
            for req in PATIENT_REQUIRED:
                if not doc.get(req):
                    errors.append({"row": idx, "field": req, "message": "Campo requerido vacío"})
                    raise ValueError("missing required")
            # Dedup
            nat = (doc.get("national_id") or "").strip().lower()
            comp = f"{doc['first_name'].strip().lower()}|{doc['last_name'].strip().lower()}|{(doc.get('phone') or '').strip()}"
            if nat and nat in existing_nat:
                skipped_dup += 1
                continue
            if comp in existing_composite:
                skipped_dup += 1
                continue
            if nat:
                existing_nat.add(nat)
            existing_composite.add(comp)

            doc["id"] = str(uuid.uuid4())
            doc["clinic_id"] = clinic_id
            doc["created_at"] = now_iso()
            doc["updated_at"] = now_iso()
            valid_rows.append(doc)
            if len(preview) < 5:
                preview.append({
                    "first_name": doc["first_name"],
                    "last_name": doc["last_name"],
                    "phone": doc.get("phone"),
                    "email": doc.get("email"),
                    "national_id": doc.get("national_id"),
                })
        except ValueError:
            continue
        except Exception as e:
            errors.append({"row": idx, "field": "*", "message": str(e)[:120]})

    imported = 0
    commit_errors = []
    if commit and valid_rows:
        BATCH = 500
        for i in range(0, len(valid_rows), BATCH):
            chunk = valid_rows[i:i + BATCH]
            try:
                sdb.table('patients').insert(chunk).execute()
                imported += len(chunk)
            except Exception as e:
                commit_errors.append({"batch": i // BATCH + 1, "message": str(e)[:200]})

    return {
        "total": len(valid_rows) + len(errors) + skipped_dup,
        "valid_rows": len(valid_rows),
        "error_rows": len(errors),
        "duplicates_skipped": skipped_dup,
        "errors": errors[:50],
        "preview": preview,
        "imported": imported,
        "committed": commit and not commit_errors,
        "commit_errors": commit_errors,
    }

@router.get("/clinic/patients-bulk/template")
async def patients_import_template(format: str = "csv", ctx=Depends(require_clinic_member)):
    """Download a CSV or XLSX template with required + optional columns and one example row."""
    headers = PATIENT_REQUIRED + PATIENT_OPTIONAL
    example = {
        "first_name": "Juan",
        "last_name": "Pérez",
        "date_of_birth": "1985-04-12",
        "gender": "male",
        "national_id": "1234567890101",
        "phone": "50255550100",
        "email": "juan.perez@ejemplo.com",
        "city": "Guatemala",
        "country": "Guatemala",
        "blood_type": "O+",
    }
    if format == "xlsx":
        from openpyxl import Workbook
        from fastapi.responses import StreamingResponse
        import io as _io
        wb = Workbook()
        ws = wb.active
        ws.title = "Pacientes"
        ws.append(headers)
        ws.append([example.get(h, "") for h in headers])
        buf = _io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return StreamingResponse(
            buf,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": 'attachment; filename="plantilla_pacientes.xlsx"'},
        )
    # CSV (default)
    csv_text = ",".join(headers) + "\n" + ",".join(f'"{example.get(h, "")}"' for h in headers) + "\n"
    return {"filename": "plantilla_pacientes.csv", "headers": headers, "content": csv_text}
